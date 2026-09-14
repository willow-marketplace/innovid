#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1"]
# ///
"""
List a workbook's sheets so the skill can offer real tab names.

Asking someone to type a sheet name means they have to open the file,
copy it exactly, and get the punctuation right — for a tab called
"Budget vs Actuals (Dept View)" that is a coin flip. Reading the tabs and
presenting them as options removes the guess.

Each sheet is reported with a shape hint so the prompt can suggest which
adapter fits, and with whether it is hidden — workbooks carry scratch and
archive tabs the author never looks at, and those are rarely the budget.

Emits JSON on stdout:

  {
    "filename": "...",
    "sheets": [
      {"name": "...", "hidden": false, "rows": 151, "cols": 64,
       "likely_shape": "tag-crosstab" | "pnl-outline" | "monthly-crosstab"
                        | "coa-mapping" | null,
       "is_mapping": false,
       "shape_settled": true,
       "why": "department super-headers over Actual/Budget blocks"}
    ]
  }

`likely_shape` orders the options and, where `shape_settled` is true, is
acted on without a second question — the header stated something a
differently-shaped sheet does not. Where it is false the classification is
a best guess and the skill still asks: a wrong guess presented as a guess
costs a glance; a wrong guess applied silently costs a wrong dashboard.

"coa-mapping" is not one of parse_budget_workbook.py's --shape values — it
names a GL-code/budget-category lookup tab, read by parse_coa_mapping.py
instead. `is_mapping` marks that case so a caller piping `likely_shape`
into --shape doesn't have to special-case the one string that isn't one.
"""

from __future__ import annotations

import argparse
import calendar
import json
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

# Quarter labels in either digit order ("Q1 2026", "1Q25", "1Q25P" for
# "Plan"), 2- or 4-digit year.
_QUARTER_RX = re.compile(
    r"^(?:Q[1-4]|[1-4]Q)[\s\-_/]*(?:\d{2}|20\d{2})[A-Za-z]*$", re.IGNORECASE
)
# A year-total column, however the workbook titles it.
_TOTAL_RX = re.compile(r"^(?:TOTAL\s+20\d{2}|20\d{2}\s+YTD)$", re.IGNORECASE)
_SUBHEADERS = {"actual", "budget", "budget variance", "% of budget"}
_SCAN_ROWS = 12
# Carta account_type codes are 4 digits, 1xxx-8xxx. A bare 3-5 digit run
# also matches invoice numbers, fee bps, and phone extensions.
_ACCOUNT_TYPE_RX = re.compile(r"\b[1-8]\d{3}\b")
# ISO date-range period labels ("2023-01-01 - 2023-01-31"); bucketed by
# the first date's month, same as any other monthly column header.
_DATE_RANGE_RX = re.compile(
    r"^(\d{4})-(\d{2})-\d{2}\s*(?:-|to|through)\s*\d{4}-\d{2}-\d{2}$",
    re.IGNORECASE,
)

# Duplicated from shapes/monthly_crosstab.py — this script runs standalone.
_MONTH_NAMES = ["jan", "feb", "mar", "apr", "may", "jun",
                "jul", "aug", "sep", "oct", "nov", "dec"]
_MONTH_FULL = ["january", "february", "march", "april", "may", "june",
               "july", "august", "september", "october", "november", "december"]
# The apostrophe in "Jan '26" is just another separator between the name
# and the year, same as the space in "Jan 26".
_MONTH_RX = re.compile(
    r"^(" + "|".join(_MONTH_NAMES + _MONTH_FULL) + r")[\s\-_/']*(\d{2}|\d{4})?$",
    re.IGNORECASE,
)


def _is_quarter_end(v: date) -> bool:
    """True when `v` is a quarter-ending month's last day — the pattern
    for date-typed quarterly columns, not twelve ordinary monthly ones."""
    return v.month in (3, 6, 9, 12) and v.day == calendar.monthrange(v.year, v.month)[1]


def _date_range_month(text: str) -> str | None:
    """3-letter month bucket for an ISO date-range label like
    "2023-01-01 - 2023-01-31", keyed off the range's first date."""
    m = _DATE_RANGE_RX.match(text.strip())
    if not m:
        return None
    return _MONTH_NAMES[int(m.group(2)) - 1]


def _parse_gl_codes(v) -> list[int]:
    if isinstance(v, int):
        return [v] if 1000 <= v <= 8999 else []
    if isinstance(v, float):
        return [int(v)] if v.is_integer() and 1000 <= v <= 8999 else []
    if isinstance(v, str):
        return [int(m) for m in _ACCOUNT_TYPE_RX.findall(v)]
    return []


def _find_gl_code_column(ws, header_row: int, sample: int = 20,
                          threshold: float = 0.6) -> int | None:
    """A column below `header_row` whose cells mostly read as short numbers.

    Same signal `parse_coa_mapping.py` uses to find the GL-code column
    independent of header wording — used here to recognize a COA-mapping
    sheet's *shape*, not to parse it (that's the other script's job once the
    sheet is picked).
    """
    max_col = min(ws.max_column, 15)
    for c in range(1, max_col + 1):
        checked = matched = 0
        for r in range(header_row + 1, min(header_row + 1 + sample, ws.max_row + 1)):
            v = ws.cell(r, c).value
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            checked += 1
            if _parse_gl_codes(v):
                matched += 1
        if checked >= 3 and (matched / checked) >= threshold:
            return c
    return None


# A GL-code column alone isn't enough evidence — plenty of sheets that
# aren't mapping tables have one (P&L, balance sheet, cash flow, any GL
# extract). Two more guards narrow this to an actual compact mapping table:
_COA_MAX_ROWS = 500          # a per-(GL x dept) reference table, not a ledger dump
_COA_MIN_SAMPLE = 10         # enough GL-code values to trust a ratio from
_COA_MIN_DUP_RATIO = 0.15    # the same GL recurring per dept/context is the tell


def _gl_column_repeat_ratio(ws, header_row: int, col: int) -> tuple[float, int]:
    """Fraction of `col`'s GL-like values (full column, not just the sample)
    that repeat, plus how many were found.

    A financial statement lists each account once — P&L, balance sheet, and
    cash flow sheets in a real workbook all showed a repeat ratio under 3%.
    A genuine mapping table restates the same GL once per department or
    context row it applies to — the same real workbook's actual mapping
    tabs showed roughly half their GL values repeated. That gap is the
    signal a fixed column position or header wording can't give: it doesn't
    care what the sheet calls its columns, only how the account-code column
    actually behaves.
    """
    seen = []
    for r in range(header_row + 1, ws.max_row + 1):
        codes = _parse_gl_codes(ws.cell(r, col).value)
        if codes:
            seen.append(codes[0])
    if not seen:
        return 0.0, 0
    return 1 - (len(set(seen)) / len(seen)), len(seen)


def _month_label_count(ws, scan_rows: int) -> int:
    """Distinct calendar months a header row names, at most. A month can be
    text ("Jan", "Jan '26") or a native Excel date, which several real
    workbooks use for their period row."""
    best = 0
    for r in range(1, min(scan_rows, ws.max_row) + 1):
        months = set()
        for c in range(1, min(ws.max_column, 120) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                s = v.strip()
                if _MONTH_RX.match(s):
                    months.add(s.lower()[:3])
                    continue
                bucket = _date_range_month(s)
                if bucket:
                    months.add(bucket)
            elif isinstance(v, date):
                months.add(_MONTH_NAMES[v.month - 1])
        best = max(best, len(months))
    return best


def _tag_super_header_columns(ws, subheader_row: int) -> int:
    """Columns above `subheader_row` carrying a real name rather than a
    period label, a sub-header word, or a bare number."""
    cols = set()
    for r in range(1, subheader_row):
        for c in range(1, min(ws.max_column, 80) + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            s = v.strip()
            if not s or s.lower() in _SUBHEADERS:
                continue
            if _QUARTER_RX.match(s) or _TOTAL_RX.match(s) or _MONTH_RX.match(s):
                continue
            if s.replace(",", "").replace(".", "").isdigit():
                continue
            cols.add(c)
    return len(cols)


# Which classifications rest on evidence strong enough to act on without
# asking, and which are a best guess worth confirming.
#
# A sheet is only ever classified from its header block, so "settled" means
# the header states something a differently-shaped sheet does not: four or
# more quarter columns, twelve months plus a year total, or named
# super-headers over repeated Actual/Budget blocks. Those are structural.
#
# The two guesses below them are weaker by construction. `monthly-crosstab`
# fires on twelve month labels plus a single Actual/Budget marker, which a
# monthly outline also satisfies; it has never been validated against a real
# customer workbook. `coa-mapping` rests on a repeat ratio tuned on one
# firm's tabs. Both stay askable.
_SETTLED_REASONS = (
    "quarterly period columns",
    "monthly period columns with a year total",
    "named super-headers over repeated Actual/Budget/Variance blocks",
)


def shape_is_settled(shape, why):
    """Whether a classification is evidence, or a guess worth confirming.

    Read by the skill to decide between stating the shape and asking about
    it. Kept beside `_classify` so the two cannot drift: a new branch is
    unsettled until it is named here deliberately.
    """
    if not shape:
        return False
    return any(str(why or "").startswith(r) for r in _SETTLED_REASONS)


def _classify(ws) -> tuple[str | None, str]:
    """Guess which adapter suits a sheet, from its header block alone.
    A "coa-mapping" guess is not a valid --shape; main() flags that below."""
    subheader_hits = 0
    first_subheader_row = None
    quarters = 0
    total_period = False
    for r in range(1, min(_SCAN_ROWS, ws.max_row) + 1):
        for c in range(1, min(ws.max_column, 80) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                s = v.strip().lower()
                if s in _SUBHEADERS:
                    subheader_hits += 1
                    if first_subheader_row is None:
                        first_subheader_row = r
                if _QUARTER_RX.match(v.strip()):
                    quarters += 1
                elif _TOTAL_RX.match(v.strip()):
                    total_period = True
            elif isinstance(v, date) and _is_quarter_end(v):
                quarters += 1

    # Outline first: pnl_outline.py sums the period columns itself when a
    # sheet has no year-total column, so one isn't required here either.
    if quarters >= 4:
        why = "quarterly period columns"
        return "pnl-outline", why + (" with a year total" if total_period else "")
    month_labels = _month_label_count(ws, _SCAN_ROWS)
    if month_labels >= 12 and total_period:
        return "pnl-outline", "monthly period columns with a year total"
    # Actual/Budget/Variance words alone also fire on a flat P&L that
    # repeats them per period; require named super-headers above them too.
    if (subheader_hits >= 6 and first_subheader_row
            and _tag_super_header_columns(ws, first_subheader_row) >= 2):
        return "tag-crosstab", "named super-headers over repeated Actual/Budget/Variance blocks"
    if _month_label_count(ws, _SCAN_ROWS + 8) >= 12 and subheader_hits >= 1:
        return "monthly-crosstab", "twelve month-label columns with Actual/Budget markers"

    # Neither budget shape: check for a Carta GL ↔ Budget Category mapping
    # sheet instead. A GL-code column paired with descriptive text columns
    # is necessary but not sufficient — see the repeat-ratio and row-count
    # guards above for why a plain GL-code column alone over-matches every
    # financial statement and raw ledger extract in the workbook.
    if ws.max_row <= _COA_MAX_ROWS:
        for r in range(1, min(_SCAN_ROWS, ws.max_row) + 1):
            col = _find_gl_code_column(ws, r)
            if col is None:
                continue
            ratio, n = _gl_column_repeat_ratio(ws, r, col)
            if n >= _COA_MIN_SAMPLE and ratio >= _COA_MIN_DUP_RATIO:
                return ("coa-mapping",
                        "a GL-code column repeating per department/context, "
                        "paired with descriptive text columns")

    return None, "no recognizable budget or mapping header block"


def find_by_sheets(directory: Path, required: list[str], limit: int = 5) -> list[dict]:
    """Workbooks in `directory` that carry every sheet in `required`.

    For finding a workbook that moved. A budget ref names both the file and
    the tabs it read, and while the path stops being true the moment
    someone renames the file, the tab names travel with it — so a workbook
    holding exactly those tabs is almost certainly the same workbook under
    a new name.

    Ranked by how much of the sheet list matched, then by recency, so a
    genuine duplicate loses to the file actually being worked on. Never
    decides: the caller shows the candidates and asks.
    """
    want = {s.strip().lower() for s in required if s and s.strip()}
    if not want:
        return []
    out = []
    for p in sorted(directory.glob("*.xlsx")):
        # Excel's lock files are not workbooks and cannot be opened.
        if p.name.startswith("~$"):
            continue
        try:
            wb = load_workbook(p, read_only=True)
            names = {n.strip().lower() for n in wb.sheetnames}
            wb.close()
        except Exception:
            continue
        hit = want & names
        if not hit:
            continue
        out.append({
            "path":     str(p),
            "filename": p.name,
            "matched":  sorted(hit),
            "missing":  sorted(want - names),
            "complete": len(hit) == len(want),
            "mtime":    p.stat().st_mtime,
        })
    out.sort(key=lambda c: (not c["complete"], -len(c["matched"]), -c["mtime"]))
    return out[:limit]


# What a sheet's rows turn out to be, and how sure that is.
#
# Two shapes read the same layout — twelve months of Budget/Actual/Variance
# over a column of labels — and differ only in what those labels ARE. The
# header block cannot tell them apart, because it is identical. The labels
# can: on one real firm's pair of tabs, 35 of 36 rows named a Carta account
# on the tab whose rows are accounts, and 6 of 20 on the tab whose rows are
# their own buckets.
#
# Claiming "these are Carta accounts" wants strong evidence, because the two
# mistakes are not equal. Claim it wrongly and the flat read finds almost
# nothing and refuses outright. Miss it and the sheet reads as an outline,
# where every row that IS an account still resolves by name — the forgiving
# direction.
#
# Measured across the candidate budget sheets of three real firms: 92%, 24%,
# 2%, and one at 62%. The first three are decisive; the fourth is a sheet
# that lists some accounts by name and some of its own headings, and no
# threshold can read it correctly without asking. So there are three
# verdicts, not two — the middle band reports `unclear` rather than picking
# the nearer edge.
_ROW_AXIS_MIN = 0.85
_ROW_AXIS_MAX_OWN = 0.5

# Below this many labelled rows the share means little either way.
_ROW_AXIS_MIN_ROWS = 4

_ROW_AXIS_SKIP_RX = re.compile(r"^(total|net|subtotal)\b", re.IGNORECASE)


def read_account_names(path):
    """{gl: name} from the chart-of-accounts dump Step 3 saves.

    A pipe table: `ACCT_TYPE | ACCOUNT` under a header line. Read here
    rather than imported so this stays a standalone inspector.
    """
    out = {}
    try:
        text = Path(path).expanduser().read_text()
    except Exception:
        return out
    for line in text.splitlines():
        if "|" not in line:
            continue
        parts = [p.strip() for p in line.split("|")]
        if len(parts) < 2:
            continue
        try:
            gl = int(parts[0])
        except ValueError:
            continue
        if parts[1]:
            out.setdefault(gl, parts[1])
    return out


def _norm_name(v):
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(v or "").lower()).split())


def _first_data_row(ws, scan=_SCAN_ROWS):
    """The row after the header block, so the title band is not counted.

    Anchored on the Actual/Budget/Variance sub-header every one of these
    sheets carries. Without it a four-line title band and a pair of section
    words dilute the share enough to flip the answer: one real tab reads 35
    of 43 scanned rows, and 35 of 36 real ones.
    """
    for r in range(1, min(scan, ws.max_row) + 1):
        for c in range(1, min(ws.max_column, 80) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().lower() in _SUBHEADERS:
                return r + 1
    return 1


def _row_axis(ws, account_names, scan=400):
    """Whether this sheet's rows are Carta's accounts or the firm's own.

    Reads the label column rather than the header, because the header is
    what the two shapes have in common.
    """
    if not account_names:
        return None
    known = {_norm_name(n) for n in account_names.values() if _norm_name(n)}
    first = _first_data_row(ws)
    best = None
    for col in (1, 2):
        labels = []
        for r in range(first, min(first + scan, ws.max_row) + 1):
            v = ws.cell(r, col).value
            if not isinstance(v, str):
                continue
            t = v.strip()
            if not t or _ROW_AXIS_SKIP_RX.match(t):
                continue
            labels.append(t)
        if len(labels) < _ROW_AXIS_MIN_ROWS:
            continue
        hits = sum(1 for t in labels if _norm_name(t) in known)
        share = hits / len(labels)
        if best is None or share > best["share"]:
            best = {"column": col, "rows": len(labels),
                    "matched": hits, "share": round(share, 3)}
    if best is None:
        return None
    if best["share"] >= _ROW_AXIS_MIN:
        best["verdict"] = "carta-accounts"
    elif best["share"] <= _ROW_AXIS_MAX_OWN:
        best["verdict"] = "own-categories"
    else:
        best["verdict"] = "unclear"
    return best


def main() -> int:
    ap = argparse.ArgumentParser(
        description="List a workbook's sheets, with a shape hint per sheet."
    )
    ap.add_argument("--workbook", help="Path to the .xlsx file")
    ap.add_argument(
        "--find-in",
        help=(
            "Directory to search for a workbook carrying --require-sheets. "
            "Used to recover a budget whose file was renamed or moved."
        ),
    )
    ap.add_argument(
        "--require-sheets",
        nargs="+",
        default=[],
        help="Sheet names the workbook must carry, for --find-in.",
    )
    ap.add_argument(
        "--accounts",
        help=(
            "Path to the firm's chart-of-accounts dump (accounts-all.txt). "
            "Adds a row_axis per sheet: whether its rows name Carta's own "
            "accounts or the firm's own categories. Omitted, row_axis is "
            "null and nothing is inferred from it."
        ),
    )
    args = ap.parse_args()

    if args.find_in:
        d = Path(args.find_in).expanduser().resolve()
        if not d.is_dir():
            print(json.dumps({"error": f"not a directory: {d}"}))
            return 2
        print(json.dumps({
            "searched":   str(d),
            "candidates": find_by_sheets(d, args.require_sheets),
        }, indent=2))
        return 0

    if not args.workbook:
        print(json.dumps({"error": "pass --workbook, or --find-in with --require-sheets"}))
        return 2

    path = Path(args.workbook).expanduser().resolve()
    if not path.exists():
        print(json.dumps({"error": f"workbook not found: {path}"}))
        return 2
    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        print(json.dumps({"error": f"could not open workbook: {e}"}))
        return 3

    account_names = read_account_names(args.accounts) if args.accounts else {}

    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        shape, why = _classify(ws)
        sheets.append({
            "name":         name,
            "hidden":       ws.sheet_state != "visible",
            "rows":         ws.max_row,
            "cols":         ws.max_column,
            "likely_shape": shape,
            # "coa-mapping" is a real hint but not a valid --shape value.
            "is_mapping":   shape == "coa-mapping",
            "why":          why,
            # Whether the classification is structural evidence or a best
            # guess — see `shape_is_settled`. The skill states a settled
            # shape and asks about an unsettled one.
            "shape_settled": shape_is_settled(shape, why),
            "row_axis":     _row_axis(ws, account_names),
        })

    # Sheets that look like budgets first, then by size — the budget tab is
    # rarely the last one in a financial-statements workbook.
    sheets.sort(key=lambda s: (s["hidden"], s["likely_shape"] is None, -s["rows"]))
    print(json.dumps({"filename": path.name, "path": str(path), "sheets": sheets}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
